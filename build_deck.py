"""Step 3: build the PowerPoint write-up from the analysis tables and rendered figures.

Run build_visuals.py first (it renders output/figures/*.png), then:
    python build_deck.py
Output: output/Netflix_Music_Project_Report.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from src import analysis
from src import content as C

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
FIG = OUT / "figures"

RED = RGBColor(0xE5, 0x09, 0x14)
INK = RGBColor(0x1A, 0x1A, 0x1A)
INK2 = RGBColor(0x5F, 0x5F, 0x5F)
PANEL = RGBColor(0xF5, 0xF5, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x14, 0x14, 0x14)
FONT = "Inter"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def fig(name):
    hits = sorted(FIG.glob(f"*_{name}.png"))
    if not hits:
        raise FileNotFoundError(f"figure '{name}' not found; run build_visuals.py first")
    return hits[0]


def _text(slide, x, y, w, h, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, FONT
        p.space_after = Pt(size * 0.45)
    return tb


def _rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def base(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    _rect(s, 0, 0, 13.333, 0.11, RED)
    _text(s, 0.6, 0.45, 12.1, 0.8, title, size=28, bold=True)
    if subtitle:
        _text(s, 0.6, 1.2, 12.1, 0.5, subtitle, size=15, color=INK2)
    _text(s, 11.6, 7.0, 1.2, 0.3, str(len(prs.slides)), size=10, color=INK2, align=PP_ALIGN.RIGHT)
    return s


def image_slide(name, notes):
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(str(fig(name)), 0, 0, width=prs.slide_width, height=prs.slide_height)
    _text(s, 11.6, 7.0, 1.2, 0.3, str(len(prs.slides)), size=10, color=INK2, align=PP_ALIGN.RIGHT)
    _notes(s, notes)
    return s


def section(num, title, why):
    s = prs.slides.add_slide(BLANK)
    _rect(s, 0, 0, 13.333, 7.5, DARK)
    _rect(s, 0.6, 2.55, 0.12, 2.0, RED)
    _text(s, 0.95, 2.4, 11, 0.6, f"PART {num}", size=16, color=RED, bold=True)
    _text(s, 0.95, 2.95, 11.5, 1.0, title, size=40, color=WHITE, bold=True)
    _text(s, 0.95, 3.95, 11.5, 0.8, why, size=18, color=RGBColor(0xC8, 0xC8, 0xC8))
    return s


def bullets_slide(title, subtitle, items, notes=None, size=18):
    s = base(title, subtitle)
    _text(s, 0.8, 1.95, 11.8, 5.0, [f"•  {t}" for t in items], size=size)
    if notes:
        _notes(s, notes)
    return s


def table_slide(title, subtitle, header, rows, widths, notes=None, size=12, top=1.9):
    s = base(title, subtitle)
    h = min(5.2, 0.42 * (len(rows) + 1) + 0.2 * len(rows))
    shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.6), Inches(top), Inches(sum(widths)), Inches(h))
    tbl = shape.table
    for k, w in enumerate(widths):
        tbl.columns[k].width = Inches(w)
    for r, row in enumerate([header] + rows):
        for k, val in enumerate(row):
            cell = tbl.cell(r, k)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RED if r == 0 else (PANEL if r % 2 else WHITE)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.name = FONT
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else INK
            cell.margin_left = cell.margin_right = Inches(0.08)
    if notes:
        _notes(s, notes)
    return s


# ---- data ------------------------------------------------------------------------
users, journey, features = analysis.load()
scores = analysis.score_features(features)

# ---- 1. title and framing -----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, 13.333, 7.5, DARK)
_rect(s, 0, 0, 13.333, 0.11, RED)
_text(s, 0.8, 1.8, 11.5, 0.6, "PRODUCT MANAGEMENT", size=16, color=RED, bold=True)
_text(s, 0.8, 2.45, 11.5, 1.6, "Launching Netflix Music", size=54, color=WHITE, bold=True)
_text(s, 0.8, 3.55, 11.5, 0.8, "Product sense, problem definition and MVP for a personalized audio app",
      size=22, color=RGBColor(0xC8, 0xC8, 0xC8))
_rect(s, 0.8, 5.05, 0.6, 0.04, RED)
_text(s, 0.8, 5.25, 11.5, 0.5, "Sharmistha Chakrabortty", size=20, color=WHITE, bold=True)
_text(s, 0.8, 5.7, 11.5, 0.4, "Aspiring Product Manager", size=15, color=RGBColor(0xC8, 0xC8, 0xC8))
_text(s, 0.8, 6.3, 11.5, 0.4, "UC San Diego Extended Studies  |  AI-Powered Product Management Professional Program",
      size=14, color=RGBColor(0x9A, 0x9A, 0x9A))

bullets_slide("Agenda", "The story runs from why, to who, to what we build first", [
    "Part 1  The art of why: why Netflix, why now, the business case",
    "Part 2  Defining the problem: customer research, personas, jobs to be done",
    "Part 3  Where the experience breaks: user journey maps",
    "Part 4  From ideas to features: Idea Funnel",
    "Part 5  Where Netflix can win: competition, SWOT, BCG",
    "Part 6  What to build first: RICE, WSJF, Kano, MoSCoW",
    "Part 7  Delivering the MVP: definition, scope, user stories, PRD",
    "Part 8  Roadmap, success metrics and risks",
])

s = base("Executive summary", "Netflix Music turns what you watch into what you listen to")
blocks = [
    ("THE PROBLEM", "Listeners can't easily find audio that fits them. 4 of 10 users report a discovery pain and "
                    "3 of 10 report missing exclusive or regional content. No app links watching and listening."),
    ("THE INSIGHT", "Netflix owns the shows, the soundtracks and the viewing data. That is an advantage "
                    "Spotify, Apple Music and Amazon Music cannot copy."),
    ("THE MVP", "Watch-history discovery and AI playlists (top 2 on RICE and WSJF), soundtrack hubs with "
                "exclusive audio, offline listening and an ad-free premium tier. Launch target: Binge-Watchers."),
    ("THE PLAN", "MVP by mid-2027, then User Growth, Competitive Differentiation and Monetization through 2028, "
                 "each phase gated by a measurable target."),
]
for i, (h, b) in enumerate(blocks):
    x = 0.6 + (i % 2) * 6.15
    y = 1.9 + (i // 2) * 2.55
    _rect(s, x, y, 5.95, 2.35, PANEL)
    _text(s, x + 0.25, y + 0.2, 5.5, 0.4, h, size=14, color=RED, bold=True)
    _text(s, x + 0.25, y + 0.65, 5.5, 1.6, b, size=15)

# ---- Part 1: art of why ---------------------------------------------------------------
section(1, "The art of why", "Before deciding what to build, be sure why it should exist and why Netflix should build it")
image_slide("five_whys", "Five whys: each answer becomes the next question until we reach something only Netflix can do. "
                         "The root answer (owned shows, soundtracks and viewing data) is what the MVP is built on.")
s = base("Business case: why Netflix should build this", "Turn unmonetized listening hours into revenue and loyalty")
_text(s, 0.6, 1.9, 6.0, 0.4, "WHY BUILD IT", size=14, color=RED, bold=True)
_text(s, 0.6, 2.35, 6.0, 4.5, [
    "•  Capture the non-screen hours: commuting, work, exercise",
    "•  Defend the core subscription: a second daily habit lowers churn",
    "•  Monetize owned IP twice: soundtracks, scores, cast podcasts, commentaries",
    "•  Feed the recommendation engine in both directions: video and audio",
], size=16)
_text(s, 6.9, 1.9, 6.0, 0.4, "BENEFITS", size=14, color=RED, bold=True)
_text(s, 6.9, 2.35, 6.0, 4.5, [
    "Users: music matched to their shows; easier discovery; exclusives; regional content; offline, ad-free listening",
    "Netflix: new premium and bundle revenue; lower churn; longer engagement; stronger regional growth; richer data",
    "Revenue model: free tier with ads, premium add-on, Netflix + Music bundles, later hi-fi and partnerships",
], size=15)
_notes(s, "Business case framed around the art of why: the product must earn its place for users and for Netflix.")
image_slide("problem_definition", "The problem statement: who, what, evidence from the dataset, why it matters, and the "
                                  "How Might We question that frames the whole product.")
table_slide("Project scope: six steps, each answering a question", "What the problem statement asks for and why",
            ["Step", "What we do", "Data", "Question it answers"],
            [["1. Personas", "Segment users, build personas", "Sheet 1", "Who are we building for?"],
             ["2. Journey maps", "Map discovery to daily use", "Sheet 2", "Where does the experience break?"],
             ["3. Idea Funnel", "Generate, screen and refine ideas", "Sheets 1-2", "What could solve it?"],
             ["4. SWOT and BCG", "Benchmark Spotify, Apple, Amazon", "Market", "Where can Netflix win?"],
             ["5. Prioritization", "RICE, WSJF, Kano, MoSCoW", "Sheet 3", "What do we build first?"],
             ["6. Roadmap", "MVP, Growth, Differentiation, Monetization", "All", "When and how do we deliver?"]],
            [2.2, 4.3, 1.6, 4.0], size=14)

# ---- Part 2: customer research --------------------------------------------------------
section(2, "Defining the problem", "Who the users are, what hurts, and what they would pay for (Sheet 1)")
image_slide("segment_overview", "Three segments: Binge-Watchers (4), Casual Listeners (3), Regional Listeners (3). "
                                "Binge-Watchers have no low tech-comfort users, so they are the natural early adopters.")
image_slide("age_distribution", "Users are 22 to 45, average 31.4. Regional Listeners are the youngest (27.3).")
image_slide("pain_point_themes", "Discovery and personalization is the most shared pain (4 of 10), followed by content "
                                 "and language gaps (3 of 10). This is why AI discovery leads the MVP.")
image_slide("purchase_motivations", "Exclusive content and a better listening experience tie as reasons to pay, so premium "
                                    "bundles both: exclusives plus ad-free, offline and quality.")
for p in C.PERSONAS:
    image_slide(f"persona_{p['segment'].split()[0].lower()}", f"Persona built from {p['based_on']}.")
image_slide("jobs_to_be_done", "Jobs to be done: the job each segment hires Netflix Music for. Product sense means "
                               "designing for the job, not the feature list.")

# ---- Part 3: journeys ------------------------------------------------------------------
section(3, "Where the experience breaks", "End-to-end journey for each persona (Sheet 2)")
for seg in C.JOURNEYS:
    image_slide(f"journey_{seg.split()[0].lower()}", f"Journey map for the {seg}. Behaviour row from Sheet 2; the other "
                                                     "rows are inferred from Sheets 1 and 2. Rebuild in Figma as one frame per persona.")
image_slide("journey_insights", "Cross-journey insights that shape the MVP and roadmap.")

# ---- Part 4: ideas ----------------------------------------------------------------------
section(4, "From ideas to features", "Idea Funnel: generate widely, then filter hard")
image_slide("idea_funnel", "20 ideas, 15 pass the user-need screen, 12 after merging overlaps, 5 selected for the MVP.")
image_slide("idea_list", "Each idea traced to a need in the dataset and its outcome.")

# ---- Part 5: competition ------------------------------------------------------------------
section(5, "Where Netflix can win", "Industry trends, competitor benchmark, SWOT and BCG")
bullets_slide("Industry trends shaping the launch", "Qualitative market view as of 2026", [
    "Audio is converging: music, podcasts and audiobooks in one app",
    "AI personalization is now expected, not novel",
    "Video and audio are blending: video podcasts and clips favour a video company",
    "Bundling drives retention: Apple One and Amazon Prime show the model",
    "Regional growth: new subscribers increasingly come from non-English markets",
])
image_slide("competitor_benchmark", "Netflix Music would lead on exclusive content but trail on catalogue size; "
                                    "ratings are analyst judgement.")
image_slide("swot", "SO strategy: use owned franchises and viewing data to attack discovery and exclusive-content gaps. "
                    "Main risk is licensing, so lean on audio Netflix already controls.")
image_slide("bcg_matrix", "Netflix Music is a Question Mark. Video streaming is the Cash Cow that funds it; exclusive "
                          "audio is the potential Star.")

# ---- Part 6: prioritization ---------------------------------------------------------------
section(6, "What to build first", "Four frameworks applied to the five features in Sheet 3")
image_slide("rice", "RICE = Reach x Impact x Confidence / Effort. Top two score about double the third.")
image_slide("wsjf", "WSJF = (Business Value + Time Criticality + Risk Reduction) / Effort. Top two tie at 4.8.")
image_slide("value_effort", "The top two sit in the quick-win zone: highest value, lowest effort (5).")
image_slide("kano", "Kano: watch-history discovery is the delighter; AI and mood playlists are performance features; "
                    "offline, sync and basic sharing are must-be.")
table_slide("MoSCoW: the launch scope", "Combining scores with product judgement",
            ["Priority", "Feature", "Why"],
            [[C.MOSCOW[f], f, why] for f, why in [
                ("AI-Generated Personalized Playlists", "Top RICE, tied top WSJF; solves the most common pain"),
                ("Music Discovery Based on Netflix Watch History", "Core differentiator; tied top WSJF"),
                ("Mood-Based Music Recommendations", "Third on both; key for Casual Listeners"),
                ("Social Sharing and Collaborative Playlists", "Mid scores; share link is cheap, collaboration can wait"),
                ("Podcast + Music Hybrid Experience", "Lowest scores, highest effort; revisit later")]],
            [2.2, 4.6, 5.3], size=14)
image_slide("combined_ranking", "All four frameworks agree on the order, so the decision is robust.")

# ---- Part 7: MVP ------------------------------------------------------------------------
section(7, "Delivering the MVP", "The smallest product that proves Netflix can win in audio")
image_slide("mvp_definition", "Each MVP feature maps to a defined problem, a hypothesis to test, and the metric "
                              "that proves or disproves it.")
image_slide("mvp_scope", "Scope discipline: build the advantage first, defer everything else on purpose.")
table_slide("User stories for the MVP and first growth phase", "As a [persona], I want [goal], so that [benefit]",
            ["ID", "As a", "I want", "So that", "Priority"],
            [list(u) for u in C.USER_STORIES], [0.9, 2.0, 4.1, 3.9, 1.1], size=11, top=1.75)
table_slide("PRD summary: Netflix Music MVP", "Mobile-first, Netflix sign-in, free with ads plus ad-free premium",
            ["Area", "Requirement"],
            [["Goal", "Prove Netflix viewers adopt and keep listening; convert through exclusives and ad-free"],
             ["Target users", "Primary: Binge-Watchers. Secondary: Casual Listeners. Regional from Phase 3"],
             ["P0 features", "Netflix sign-in; watch-history recs; 'Listen' after credits; soundtrack hubs; AI daily mix; "
                             "search; offline downloads; free and premium tiers"],
             ["P1 features", "Share link; mood tiles (beta)"],
             ["Non-functional", "Playback < 2 s on 4G; 99.9% availability; opt-in use of viewing history; accessible"],
             ["Dependencies", "Label licensing; exclusive audio for launch titles; Netflix rec platform and accounts"],
             ["Release criteria", "P0 stories pass; 40% of beta users listen weekly; no open privacy or licensing issue"],
             ["Non-goals", "Full catalogue parity; podcast hybrid feed; collaborative playlists; wide regional rollout"]],
            [2.4, 9.7], size=13)

# ---- Part 8: roadmap ---------------------------------------------------------------------
section(8, "Roadmap and success", "From MVP to monetization, with a gate for every phase")
image_slide("roadmap", "Four phases from January 2027 to December 2028; each must pass its gate before scaling.")
image_slide("kpis", "Year-one targets are planning assumptions to validate in the closed beta.")
table_slide("Key risks and how we reduce them", None, ["Risk", "Mitigation"], [list(r) for r in C.RISKS],
            [4.0, 8.1], size=15)

s = base("Conclusion", "Start with why, define the problem, deliver the smallest product that proves it")
_text(s, 0.6, 1.9, 12, 4.8, [
    "•  Why: listening hours are unmonetized, and Netflix owns an advantage rivals cannot copy",
    "•  Problem: no app connects what people watch with what they hear; discovery and exclusives are the gaps",
    "•  Who: Binge-Watchers first, then Casual and Regional Listeners",
    "•  MVP: watch-history discovery, AI playlists, soundtrack hubs, offline, ad-free premium",
    "•  Proof: 40% weekly active, 30% Netflix-linked listening, 25% conversion",
    "•  Next: grow, differentiate regionally, then monetize at scale through 2028",
], size=19)
_text(s, 0.6, 6.4, 12, 0.5, "“How might we turn what members watch into what they love to listen to?”",
      size=18, color=RED, bold=True)

OUT.mkdir(exist_ok=True)
dest = OUT / "Netflix_Music_Project_Report.pptx"
prs.save(dest)
print(f"Wrote {dest} ({len(prs.slides)} slides)")
